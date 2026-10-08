#!/usr/bin/env python3
"""Собирает Word-документ программы концерта из Markdown.

    python3 tools/build-program-doc.py <slug>

Источник: projects/<slug>/output/program.md
Результат: projects/<slug>/output/program.docx

Оформление по brand/BRAND.md: красный #ff2626 только в линиях и крупных цифрах,
текст #2e2e2e, приглушённый серый не светлее #6f6f6f, фон белый. Шрифт один — Arial:
дисплейный Bebas в Word не вложить, его роль играет Arial Bold прописными.
Пометки [УТОЧНИТЬ] набираются тёмным жирным на светло-сером: мелкий текст
красным не набирается.

Что понимает разметка: «# Название», абзац-подзаголовок сразу под ним,
строки «**Ключ:** значение» подряд — таблица фактов, «## Раздел», нумерованные
пункты — строки программы, «- пункт» — список, остальное — абзацы.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
INK, RED, MUTED = RGBColor(0x2E, 0x2E, 0x2E), RGBColor(0xFF, 0x26, 0x26), RGBColor(0x6F, 0x6F, 0x6F)
SURFACE, RULE = "F2F2F2", "E8E8E8"
TOKEN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\(https?://[^)]+\)|\[УТОЧНИТЬ[^\]]*\])")


def set_font(run, size, color=INK, bold=False, caps=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rpr.rFonts.set(qn(attr), FONT)
    run.font.size, run.font.color.rgb = Pt(size), color
    run.font.bold, run.font.all_caps = bold, caps


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def border(parent, side, color, sz):
    el = OxmlElement(f"w:{side}")
    for k, v in (("val", "single" if sz else "nil"), ("sz", str(sz)), ("space", "0"), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    parent.append(el)


def para_border(p, side, color, sz, space=4):
    ppr = p._p.get_or_add_pPr()
    box = ppr.find(qn("w:pBdr"))
    if box is None:
        box = OxmlElement("w:pBdr")
        ppr.append(box)
    el = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", str(space)), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    box.append(el)


def spacing(p, before=0, after=6, line=1.25):
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line


def link(p, url, text, size):
    """Ссылка цветом текста, красное — подчёркивание (правило бренда)."""
    rid = p.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        f.set(qn(f"w:{attr}"), FONT)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "2E2E2E")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    u.set(qn("w:color"), "FF2626")
    for e in (f, c, sz, u):
        pr.append(e)
    r.append(pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def inline(p, text, size=10.5, color=INK, bold=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        m = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", part)
        if m:
            link(p, m.group(2), m.group(1), size)
        elif part.startswith("[УТОЧНИТЬ"):
            r = p.add_run(part)
            set_font(r, size, INK, bold=True)
            shade(r._element.get_or_add_rPr(), SURFACE)
        elif part.startswith("**"):
            set_font(p.add_run(part[2:-2]), size, color, bold=True)
        else:
            set_font(p.add_run(part), size, color, bold=bold)


def cell_margins(table, top=90, bottom=90, left=0, right=120):
    pr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    pr.append(m)


def table(doc, widths, rules=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    cell_margins(t)
    t._widths = widths
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)
    for col, w in zip(t.columns, widths):
        col.width = w
    return t


def row(t, rule=RULE):
    r = t.add_row()
    for c, w in zip(r.cells, t._widths):
        c.width = w
        tcb = OxmlElement("w:tcBorders")
        for side in ("top", "left", "right"):
            border(tcb, side, "FFFFFF", 0)
        border(tcb, "bottom", rule, 4)
        c._tc.get_or_add_tcPr().append(tcb)
    return r


def cell_text(c, text, size=10.5, color=INK, bold=False, caps=False):
    p = c.paragraphs[0]
    spacing(p, 0, 0, 1.2)
    if caps:
        set_font(p.add_run(text), size, color, bold=bold, caps=True)
    else:
        inline(p, text, size, color, bold)
    return p


def build(md: str) -> Document:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin, s.bottom_margin = Cm(1.8), Cm(2.0)
    doc.styles["Normal"].font.name = FONT
    width = Cm(16.6)

    # верхняя планка: имя лейбла и вид документа, под ней жирная красная линия
    p = doc.add_paragraph()
    spacing(p, 0, 0, 1.0)
    set_font(p.add_run("FANCYMUSIC"), 9, INK, bold=True, caps=True)
    set_font(p.add_run("   ·   Программа концерта"), 9, MUTED, caps=True)
    para_border(p, "bottom", "FF2626", 36, space=8)

    lines = [l for l in md.split("\n") if l.strip()]
    i, first_para, facts = 0, True, None
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            p = doc.add_paragraph()
            spacing(p, 26, 8, 0.95)
            set_font(p.add_run(ln[2:]), 32, INK, bold=True, caps=True)
        elif ln.startswith("## "):
            p = doc.add_paragraph()
            spacing(p, 24, 8, 1.0)
            p.paragraph_format.keep_with_next = True
            title = ln[3:]
            set_font(p.add_run(title), 15, INK, bold=True, caps=True)
            para_border(p, "bottom", "2E2E2E", 8, space=5)
            facts = None
        elif re.match(r"^\*\*[^*]+:\*\*", ln):
            if facts is None:
                spacing(doc.add_paragraph(), 0, 4, 1.0)
                facts = table(doc, [Cm(4.6), Cm(12.0)])
            k, v = re.match(r"^\*\*([^*]+):\*\*\s*(.*)$", ln).groups()
            r = row(facts)
            cell_text(r.cells[0], k, 8.5, MUTED, bold=True, caps=True)
            cell_text(r.cells[1], v, 11)
        elif re.match(r"^\d+\. ", ln):
            t = table(doc, [Cm(1.4), Cm(11.0), Cm(4.2)])
            while i < len(lines) and re.match(r"^\d+\. ", lines[i]):
                num, body = re.match(r"^(\d+)\. (.*)$", lines[i]).groups()
                body, _, dur = body.partition(" · ")
                r = row(t)
                p = r.cells[0].paragraphs[0]
                spacing(p, 0, 0, 1.0)
                set_font(p.add_run(num), 26, RED, bold=True)   # крупный красный: от 24 pt
                cell_text(r.cells[1], body, 11)
                if dur:
                    cell_text(r.cells[2], dur, 11)
                i += 1
            i -= 1
            spacing(doc.add_paragraph(), 0, 6, 1.0)
            facts = None
        elif ln.startswith("- "):
            p = doc.add_paragraph()
            spacing(p, 0, 5, 1.25)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            set_font(p.add_run("— "), 10.5, INK)
            inline(p, ln[2:], 10.5)
        else:
            p = doc.add_paragraph()
            if first_para:   # подзаголовок под названием
                spacing(p, 0, 4, 1.3)
                inline(p, ln, 12, MUTED)
                first_para = False
            else:
                spacing(p, 0, 8, 1.35)
                inline(p, ln, 10.5)
        i += 1

    # низ страницы
    foot = s.footer.paragraphs[0]
    spacing(foot, 0, 0, 1.0)
    para_border(foot, "top", "E8E8E8", 6, space=6)
    set_font(foot.add_run("FANCYMUSIC · fancymusic.ru · t.me/fancymusiclabel"), 8.5, MUTED)
    doc.core_properties.title = md.split("\n", 1)[0].lstrip("# ").strip()
    return doc


def main():
    slug = sys.argv[1]
    out = ROOT / "projects" / slug / "output"
    build((out / "program.md").read_text(encoding="utf-8")).save(out / "program.docx")
    print(f"готово: {out / 'program.docx'}")


if __name__ == "__main__":
    main()
