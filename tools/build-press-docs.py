#!/usr/bin/env python3
"""Собирает DOCX и PDF пресс-релиза.

    python3 tools/build-press-docs.py <slug>

Источники ровно те же, что у страницы, и разбираются тем же кодом
(tools/release_content.py), поэтому заголовок и текст в документе
и на странице совпадают дословно:
    projects/<slug>/output/press-release.md — вся проза
    projects/<slug>/page.json — фактура: обложка, ссылки, контакты, дата

Результат:
    projects/<slug>/output/press-release.{docx,pdf}
    projects/<slug>/output/press-release.gdoc.html — разметка для импорта
        в Google Docs: те же блоки, но простой HTML, который переживает
        конвертацию. В docs/ не публикуется, нужен только для загрузки.
    копии docx и pdf рядом со страницей, в docs/releases/<slug>/

Порядок в документе повторяет страницу: обложка рядом с названием, под ними
ссылки на стриминги и на страницу релиза, дальше текст, справки, списки
ссылок, контакты. Строка «FANCYMUSIC · Пресс-релиз · дата» — в самом низу.

Обложка вставляется только из локального файла: путь берётся из
values.COVER_FILE в page.json. Ссылки на чужой домен в документ вложить
нельзя — там нужны сами байты картинки. Если файла нет, документы
собираются без неё и печатается предупреждение.

Про шрифты: Bebas и Inter в DOCX не вкладываются, у получателя их нет.
Заголовки идут Arial Narrow, текст — Arial. Фирменные цвета сохраняются.
PDF печатается через Chromium, поэтому в нём настоящие Bebas и Inter.
"""
import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import release_content

ROOT = Path(__file__).resolve().parent.parent
INK, RED, GREY = RGBColor(0x2E, 0x2E, 0x2E), RGBColor(0xFF, 0x26, 0x26), RGBColor(0x6F, 0x6F, 0x6F)
DISPLAY_FONT, TEXT_FONT = "Arial Narrow", "Arial"


# --------------------------------------------------------------------- DOCX
def style_run(run, *, font=TEXT_FONT, size=11, color=INK, bold=False, caps=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.all_caps = caps
    return run


def lead_link_html(values):
    url = values.get("LEAD_LINK_URL", "").strip()
    if not url:
        return ""
    return f'<p style="margin-top:10pt;margin-bottom:16pt"><a href="{url}">{values.get("LEAD_LINK_LABEL") or url}</a></p>'


def inline(paragraph, text, **kw):
    bold_all = kw.pop("bold", False)
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            style_run(paragraph.add_run(chunk), bold=bold_all or (i % 2 == 1), **kw)


def add_hyperlink(paragraph, url: str, text: str, *, size=10, color="FF2626"):
    """python-docx не умеет ссылки из коробки — собираем узел вручную."""
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    from docx.oxml import OxmlElement
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, attr, value in (("w:rFonts", "w:ascii", TEXT_FONT),
                             ("w:color", "w:val", color),
                             ("w:sz", "w:val", str(int(size * 2)))):
        el = OxmlElement(tag)
        el.set(qn(attr), value)
        if tag == "w:rFonts":
            el.set(qn("w:hAnsi"), value)
        props.append(el)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def build_docx(doc_data: dict, page: dict, cover: Path | None, dst: Path) -> None:
    values = page["values"]
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = TEXT_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.35

    # --- шапка: обложка слева, название справа
    header = document.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.LEFT
    header.autofit = False
    left, right = header.rows[0].cells
    left.width, right.width = Cm(4.6), Cm(11.6)

    if cover is not None:
        left.paragraphs[0].add_run().add_picture(str(cover), width=Cm(4.4))
    else:
        style_run(left.paragraphs[0].add_run("[обложка]"), size=9, color=GREY)

    title_p = right.paragraphs[0]
    title_p.paragraph_format.space_after = Pt(4)
    style_run(title_p.add_run(doc_data["title"]),
              font=DISPLAY_FONT, size=21, color=INK, bold=True, caps=True)

    # --- эпиграф сразу под заголовком, до ссылок
    if doc_data["quote"] and doc_data["quote_first"]:
        render_quote(document, *doc_data["quote"])

    # --- ссылки-кнопки под шапкой
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    for label, url in (("Ссылки на стриминги", values["PREVIEW_URL"]),
                       ("Страница релиза", values["RELEASE_URL"])):
        if not url:
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        style_run(p.add_run(f"{label}: "), size=10, color=INK, bold=True)
        add_hyperlink(p, url, url)

    # --- лид, основной текст, цитата (эпиграф — перед лидом)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    inline(document.add_paragraph(), doc_data["lead"], size=13, bold=True)
    lead_url = values.get("LEAD_LINK_URL", "").strip()
    if lead_url:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(16)
        add_hyperlink(p, lead_url, values.get("LEAD_LINK_LABEL") or lead_url, size=11)
    for block in doc_data["body"]:
        render_block(document, block)
    if doc_data["quote"] and not doc_data["quote_first"]:
        render_quote(document, *doc_data["quote"])

    add_event_docx(document, page.get("event"))

    # --- разделы из прозы
    for heading, blocks in doc_data["sections"]:
        add_heading(document, heading)
        for block in blocks:
            render_block(document, block)
        if heading in release_content.ARTIST_SECTIONS:
            add_links(document, page["values"].get("ARTIST_LINKS_LABEL", "Ссылки на страницы артиста"), page["artist_links"])
        elif heading in release_content.LABEL_SECTIONS:
            add_links(document, "Ссылки на страницы лейбла", page["label_links"])

    # --- контакты
    add_heading(document, "Контакты для прессы")
    for label, value, url in (("Email", values["PRESS_EMAIL"], None),
                              ("Telegram", values["PRESS_TELEGRAM"], None)):
        text = re.sub(r"<[^>]+>", "", value)
        href = re.search(r'href="([^"]+)"', value)
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        style_run(p.add_run(f"{label}: "), size=10, color=INK, bold=True)
        if href:
            add_hyperlink(p, href.group(1), text)
        else:
            style_run(p.add_run(text), size=10)

    # --- выходные данные внизу, как на странице
    colophon = document.add_paragraph()
    colophon.paragraph_format.space_before = Pt(20)
    style_run(colophon.add_run(f"FANCYMUSIC · Пресс-релиз · {values['RELEASE_DATE']}"),
              size=8, color=GREY, caps=True)

    document.core_properties.title = doc_data["title"]
    document.core_properties.author = "FANCYMUSIC"
    document.save(dst)


def add_heading(document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(5)
    style_run(p.add_run(text), font=DISPLAY_FONT, size=13, color=RED, bold=True, caps=True)


def flat_links(items):
    """Группы {name, links} разворачиваются в «Имя · Площадка»."""
    if items and isinstance(items[0], dict):
        return [(i, f"{g['name']} · {l}", u) for g in items for i, l, u in g["links"]]
    return items


def add_links(document, title: str, items) -> None:
    items = flat_links(items)
    if not items:
        return
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    style_run(p.add_run(title), size=8, color=GREY, caps=True)
    for _icon, label, url in items:
        row = document.add_paragraph()
        row.paragraph_format.space_after = Pt(1)
        style_run(row.add_run(f"{label} — "), size=10, color=INK)
        add_hyperlink(row, url, url)


def event_rows(ev):
    """Событие как последовательность строк — общая для DOCX, PDF и Google Docs."""
    if not ev:
        return None
    where = " · ".join(x for x in (ev.get("venue"), ev.get("address")) if x)
    rows = {"title": ev.get("title", "Презентация альбома"),
            "when": ev.get("when", ""), "where": where, "intro": ev.get("intro", ""),
            "program": ev.get("program") or [], "lineup": ev.get("lineup") or [],
            "note": ev.get("note", ""), "url": ev.get("url", ""), "price": ev.get("price", "")}
    return rows


def add_event_docx(document, ev) -> None:
    ev = event_rows(ev)
    if not ev:
        return
    add_heading(document, ev["title"])
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(ev["when"]), font=DISPLAY_FONT, size=14, color=INK, bold=True, caps=True)
    if ev["where"]:
        style_run(document.add_paragraph().add_run(ev["where"]), size=10, color=GREY)
    if ev["intro"]:
        inline(document.add_paragraph(), ev["intro"])
    for title, rows in (("Программа", ev["program"]), ("Исполнители", ev["lineup"])):
        if not rows:
            continue
        h = document.add_paragraph()
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
        style_run(h.add_run(title), size=8, color=GREY, caps=True)
        for r in rows:
            li = document.add_paragraph()
            li.paragraph_format.space_after = Pt(1)
            style_run(li.add_run(r), size=10)
    if ev["note"]:
        inline(document.add_paragraph(), ev["note"])
    if ev["url"]:
        p = document.add_paragraph()
        label = "Билеты" + (f" ({ev['price']})" if ev["price"] else "") + ": "
        style_run(p.add_run(label), size=10, color=INK, bold=True)
        add_hyperlink(p, ev["url"], ev["url"])


def event_html(ev) -> str:
    ev = event_rows(ev)
    if not ev:
        return ""
    out = [f"<h2>{ev['title']}</h2>", f"<p><b>{ev['when']}</b><br>{ev['where']}</p>"]
    if ev["intro"]:
        out.append(f"<p>{ev['intro']}</p>")
    for title, rows in (("Программа", ev["program"]), ("Исполнители", ev["lineup"])):
        if rows:
            out.append(f"<p class='small'>{title}</p><p>" + "<br>".join(rows) + "</p>")
    if ev["note"]:
        out.append(f"<p>{ev['note']}</p>")
    if ev["url"]:
        price = f" ({ev['price']})" if ev["price"] else ""
        out.append(f"<p><b>Билеты{price}:</b> <a href=\"{ev['url']}\">{ev['url']}</a></p>")
    return "\n".join(out)


def render_quote(document, body, author: str) -> None:
    verse = release_content.is_verse(body)
    for i, line in enumerate(body):
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_before = Pt(10 if i == 0 else 0)
        p.paragraph_format.space_after = Pt(2 if i == len(body) - 1 else 0)
        if verse:
            inline(p, line, size=11, color=INK)
        else:
            inline(p, line, font=DISPLAY_FONT, size=13, color=INK, caps=True)
    if author:
        a = document.add_paragraph()
        a.paragraph_format.left_indent = Cm(0.7)
        style_run(a.add_run(author), size=9, color=GREY)


def render_block(document, block: str) -> None:
    if block.startswith(">"):
        render_quote(document, *release_content.split_quote(block))
    elif all(l.lstrip().startswith(("- ", "* ")) for l in block.splitlines()):
        for line in block.splitlines():
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            inline(p, line.lstrip("-* ").strip())
    else:
        inline(document.add_paragraph(), " ".join(l.strip() for l in block.splitlines()))


# ---------------------------------------------------------------------- PDF
PRINT_CSS = """
@page { size: A4; margin: 16mm 18mm; }
body { font-family: 'FM Text', Inter, Arial, sans-serif; font-size: 10pt;
       line-height: 1.5; color: #2e2e2e; margin: 0; }
.head { display: flex; gap: 8mm; align-items: flex-start; margin-bottom: 6mm; }
.head img { width: 42mm; height: 42mm; object-fit: cover; border: .3pt solid #e8e8e8; flex: none; }
.head .slot { width: 42mm; height: 42mm; border: .5pt dashed #d4d4d4; flex: none;
              display: flex; align-items: center; justify-content: center;
              color: #9b9b9b; font-size: 8pt; }
.eyebrow { font-size: 7.5pt; letter-spacing: .12em; text-transform: uppercase;
           color: #6f6f6f; margin: 0 0 2mm; }
.eyebrow b { color: #ff2626; }
h1 { font-family: 'FM Display', Impact, sans-serif; font-weight: 700;
     text-transform: uppercase; font-size: 24pt; line-height: .95; margin: 0 0 2mm;
     letter-spacing: .01em; }
.subtitle { color: #6f6f6f; font-size: 9.5pt; margin: 0; }
.actions { margin: 0 0 6mm; }
.actions p { margin: 0 0 1mm; font-size: 9.5pt; }
.actions b { font-weight: 600; }
h2 { font-family: 'FM Display', Impact, sans-serif; font-weight: 400;
     text-transform: uppercase; font-size: 12.5pt; color: #ff2626;
     margin: 6mm 0 2mm; letter-spacing: .01em; break-after: avoid; }
p { margin: 0 0 3mm; }
.lead { font-size: 12.5pt; line-height: 1.4; }
blockquote { margin: 4mm 0; padding-left: 3mm; border-left: 2pt solid #ff2626;
             font-family: 'FM Display', Impact, sans-serif; font-weight: 400;
             text-transform: uppercase; font-size: 13pt; line-height: 1.1; }
blockquote.verse { font-family: 'FM Text', Inter, Arial, sans-serif; font-weight: 400;
                   text-transform: none; font-size: 11pt; line-height: 1.4; }
blockquote footer { font-family: 'FM Text', Inter, Arial, sans-serif;
                    text-transform: none; font-size: 8.5pt; color: #6f6f6f; margin-top: 1.5mm; }
.links { margin: 3mm 0 0; }
.links .label { font-size: 7.5pt; letter-spacing: .12em; text-transform: uppercase;
                color: #6f6f6f; margin: 0 0 1.5mm; }
.links p { margin: 0 0 1mm; font-size: 9.5pt; }
a { color: #ff2626; text-decoration: none; }
.contacts p { margin: 0 0 1mm; font-size: 9.5pt; }
.colophon { margin-top: 8mm; padding-top: 2mm; border-top: .3pt solid #e8e8e8;
            font-size: 7.5pt; letter-spacing: .12em; text-transform: uppercase;
            color: #6f6f6f; }
.colophon b { color: #ff2626; }
"""


def html_blocks(blocks) -> str:
    out = []
    for block in blocks:
        if block.startswith(">"):
            body, author = release_content.split_quote(block)
            cls = ' class="verse"' if release_content.is_verse(body) else ""
            text = "<br>".join(body)
            out.append(f"<blockquote{cls}>{text}<footer>{author}</footer></blockquote>")
        elif all(l.lstrip().startswith(("- ", "* ")) for l in block.splitlines()):
            items = "".join(f"<li>{l.lstrip('-* ').strip()}</li>" for l in block.splitlines())
            out.append(f"<ul>{items}</ul>")
        else:
            text = " ".join(l.strip() for l in block.splitlines())
            text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
            out.append(f"<p>{text}</p>")
    return "\n".join(out)


def links_html(title: str, items) -> str:
    items = flat_links(items)
    if not items:
        return ""
    rows = "".join(f'<p>{label} — <a href="{url}">{url}</a></p>' for _i, label, url in items)
    return f'<div class="links"><p class="label">{title}</p>{rows}</div>'


def build_pdf(doc_data: dict, page: dict, cover: Path | None, dst: Path) -> None:
    from playwright.sync_api import sync_playwright
    values = page["values"]

    cover_html = (f'<img src="{cover.resolve().as_uri()}" alt="">' if cover
                  else '<div class="slot">обложка</div>')
    quote_html = ""
    if doc_data["quote"]:
        body_q, author_q = doc_data["quote"]
        cls = ' class="verse"' if release_content.is_verse(body_q) else ""
        quote_html = (f"<blockquote{cls}>{'<br>'.join(body_q)}"
                      f"<footer>{author_q}</footer></blockquote>")

    sections = []
    for heading, blocks in doc_data["sections"]:
        sections.append(f"<h2>{heading}</h2>{html_blocks(blocks)}")
        if heading in release_content.ARTIST_SECTIONS:
            sections.append(links_html(page["values"].get("ARTIST_LINKS_LABEL", "Ссылки на страницы артиста"), page["artist_links"]))
        elif heading in release_content.LABEL_SECTIONS:
            sections.append(links_html("Ссылки на страницы лейбла", page["label_links"]))

    body = f"""
<div class="head">
  {cover_html}
  <div>
    <h1>{doc_data['title']}</h1>
  </div>
</div>
{quote_html if doc_data['quote_first'] else ''}
<div class="actions">
  <p><b>Ссылки на стриминги:</b> <a href="{values['PREVIEW_URL']}">{values['PREVIEW_URL']}</a></p>
  {f'<p><b>Страница релиза:</b> <a href="{values["RELEASE_URL"]}">{values["RELEASE_URL"]}</a></p>' if values['RELEASE_URL'] else ''}
</div>
<p class="lead"><b>{release_content.bold_to_html(doc_data['lead'])}</b></p>{lead_link_html(values)}
{html_blocks(doc_data['body'])}
{'' if doc_data['quote_first'] else quote_html}
{event_html(page.get('event'))}
{''.join(sections)}
<h2>Контакты для прессы</h2>
<div class="contacts">
  <p>Email — {values['PRESS_EMAIL']}</p>
  <p>Telegram — {values['PRESS_TELEGRAM']}</p>
</div>
<p class="colophon"><b>FANCYMUSIC</b> · Пресс-релиз · {values['RELEASE_DATE']}</p>
"""
    brand = (ROOT / "brand").as_uri()
    html = (f'<!doctype html><html lang="ru"><head><meta charset="utf-8">'
            f"<title>{doc_data['title']} — пресс-релиз FANCYMUSIC</title>"
            f'<link rel="stylesheet" href="{brand}/fonts.css">'
            f"<style>{PRINT_CSS}</style></head><body>{body}</body></html>")

    tmp = dst.with_suffix(".print.html")
    tmp.write_text(html)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path="/opt/pw-browsers/chromium")
            pg = browser.new_page()
            pg.goto(tmp.as_uri(), wait_until="networkidle")
            pg.wait_for_timeout(600)
            pg.pdf(path=str(dst), format="A4", print_background=True)
            browser.close()
    finally:
        tmp.unlink(missing_ok=True)


# --------------------------------------------------------------------- main
SITE_BASE = "https://skrasin.github.io/fancymusic"

GDOC_CSS = """
body { font-family: Arial, sans-serif; font-size: 11pt; color: #2e2e2e; }
h1 { font-size: 20pt; text-transform: uppercase; margin: 0 0 6pt; }
h2 { font-size: 13pt; text-transform: uppercase; color: #ff2626; margin: 18pt 0 4pt; }
blockquote { border-left: 3pt solid #ff2626; margin: 10pt 0; padding-left: 10pt;
             font-size: 11pt; }
.small { font-size: 8pt; color: #6f6f6f; text-transform: uppercase; }
"""


def build_gdoc_html(doc_data: dict, page: dict, slug: str, dst: Path) -> None:
    """Разметка для импорта в Google Docs.

    Google при конвертации выбрасывает почти весь CSS, поэтому разметка
    простая. Обложка даётся публичной ссылкой на нашу копию с GitHub Pages:
    локальный файл конвертер не видит, а картинку по http достаёт сам.
    """
    values = page["values"]
    cover_html = ""
    cover_name = values.get("COVER_FILE")
    if cover_name:
        url = f"{SITE_BASE}/releases/{slug}/{cover_name}"
        cover_html = f'<p><img src="{url}" width="260" height="260" alt=""></p>'

    quote = ""
    if doc_data["quote"]:
        body_q, author_q = doc_data["quote"]
        style = ("" if release_content.is_verse(body_q)
                 else " style='text-transform:uppercase;font-size:12pt'")
        quote = (f"<blockquote{style}>{'<br>'.join(body_q)}"
                 f"<br><span class='small'>{author_q}</span></blockquote>")

    sections = []
    for heading, blocks in doc_data["sections"]:
        sections.append(f"<h2>{heading}</h2>{html_blocks(blocks)}")
        if heading in release_content.ARTIST_SECTIONS:
            sections.append(links_html(page["values"].get("ARTIST_LINKS_LABEL", "Ссылки на страницы артиста"), page["artist_links"]))
        elif heading in release_content.LABEL_SECTIONS:
            sections.append(links_html("Ссылки на страницы лейбла", page["label_links"]))

    html = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<title>{doc_data['title']}</title><style>{GDOC_CSS}</style></head><body>
{cover_html}
<h1>{doc_data['title']}</h1>
{quote if doc_data['quote_first'] else ''}
<p><b>Ссылки на стриминги:</b> <a href="{values['PREVIEW_URL']}">{values['PREVIEW_URL']}</a>{f'<br>{chr(10)}<b>Страница релиза:</b> <a href="{values["RELEASE_URL"]}">{values["RELEASE_URL"]}</a>' if values['RELEASE_URL'] else ''}</p>
<p class="lead"><b>{release_content.bold_to_html(doc_data['lead'])}</b></p>{lead_link_html(values)}
{html_blocks(doc_data['body'])}
{'' if doc_data['quote_first'] else quote}
{event_html(page.get('event'))}
{''.join(sections)}
<h2>Контакты для прессы</h2>
<p>Email — {values['PRESS_EMAIL']}<br>Telegram — {values['PRESS_TELEGRAM']}</p>
<p class="small">FANCYMUSIC · Пресс-релиз · {values['RELEASE_DATE']}</p>
</body></html>"""
    dst.write_text(html)


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    slug = sys.argv[1]
    proj = ROOT / "projects" / slug
    src = proj / "output/press-release.md"
    page_path = proj / "page.json"
    if not src.exists() or not page_path.exists():
        print(f"нет исходников: нужны {src} и {page_path}")
        return 1

    page = json.loads(page_path.read_text())
    doc_data = release_content.parse(src.read_text())

    cover = None
    cover_name = page["values"].get("COVER_FILE")
    if cover_name:
        candidate = ROOT / "docs/releases" / slug / cover_name
        if candidate.exists():
            cover = candidate
        else:
            print(f"ВНИМАНИЕ: файл обложки не найден: {candidate}")
    else:
        print("ВНИМАНИЕ: в page.json нет values.COVER_FILE — документы собраны без обложки.")
        print("          Вложить картинку по внешней ссылке нельзя, нужен локальный файл.")

    out_dir = src.parent
    docx_path, pdf_path = out_dir / "press-release.docx", out_dir / "press-release.pdf"
    build_docx(doc_data, page, cover, docx_path)
    build_pdf(doc_data, page, cover, pdf_path)
    gdoc_path = out_dir / "press-release.gdoc.html"
    build_gdoc_html(doc_data, page, slug, gdoc_path)

    page_dir = ROOT / "docs/releases" / slug
    if page_dir.exists():
        for f in (docx_path, pdf_path):
            shutil.copy(f, page_dir / f.name)

    for f in (docx_path, pdf_path, gdoc_path):
        print(f"{f.relative_to(ROOT)} — {f.stat().st_size // 1024} КБ")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
